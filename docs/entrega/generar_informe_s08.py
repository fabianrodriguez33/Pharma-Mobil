# Genera S08_ActividadAutonoma_RodriguezBazan.docx a partir de las evidencias reales de docs/entrega/s08.
import glob
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

AQUI = os.path.dirname(os.path.abspath(__file__))
EV = os.path.join(AQUI, "s08")
TMP = os.path.join(EV, "_miniaturas")
os.makedirs(TMP, exist_ok=True)
AZUL = RGBColor(0x1F, 0x4E, 0x79)
AZUL_HEX, CELESTE_HEX, GRIS_HEX = "1F4E79", "9DC3E6", "EAF2FA"
FUENTE = "Calibri"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.2)

for nombre, tam, color in (("Normal", 10.5, None), ("Heading 1", 16, AZUL), ("Heading 2", 13, AZUL), ("Heading 3", 11, AZUL)):
    st = doc.styles[nombre]
    st.font.name = FUENTE
    st.font.size = Pt(tam)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    st.element.rPr.rFonts.set(qn("w:ascii"), FUENTE)
    st.element.rPr.rFonts.set(qn("w:hAnsi"), FUENTE)
    if color:
        st.font.color.rgb = color
        st.font.bold = True


def parrafo(texto="", negrita=False, tam=None, centrado=False, color=None, cursiva=False, espacio=4):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold, r.italic = negrita, cursiva
    if tam:
        r.font.size = Pt(tam)
    if color:
        r.font.color.rgb = color
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(espacio)
    return p


def sombrear(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def bordes(tabla_):
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + lado)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), CELESTE_HEX)
        b.append(e)
    tabla_._tbl.tblPr.append(b)


def tabla(encabezados, filas, anchos=None, tam=9):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes(t)
    for i, h in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(tam)
        r.font.color.rgb = RGBColor(255, 255, 255)
        sombrear(c, AZUL_HEX)
    for n, fila in enumerate(filas):
        cs = t.add_row().cells
        for i, v in enumerate(fila):
            cs[i].text = ""
            r = cs[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(tam)
            if n % 2 == 1:
                sombrear(cs[i], GRIS_HEX)
    if anchos:
        t.autofit = False
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def ficha(pares):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes(t)
    for k, v in pares:
        cs = t.add_row().cells
        cs[0].text = ""
        cs[1].text = ""
        r = cs[0].paragraphs[0].add_run(k)
        r.bold = True
        r.font.size = Pt(9)
        sombrear(cs[0], GRIS_HEX)
        r = cs[1].paragraphs[0].add_run(v)
        r.font.size = Pt(9)
        cs[0].width, cs[1].width = Cm(4.2), Cm(12.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def imagen(nombre, leyenda, ancho=4.3):
    im = Image.open(os.path.join(EV, nombre)).convert("RGB")
    im.thumbnail((620, 1380))
    mini = os.path.join(TMP, nombre.replace(".png", ".jpg"))
    im.save(mini, quality=82)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(mini, width=Cm(ancho))
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    parrafo(leyenda, cursiva=True, tam=8.5, centrado=True, color=RGBColor(0x59, 0x59, 0x59), espacio=8)


def codigo(texto, tam=7.5):
    for linea in texto.strip("\n").splitlines():
        p = doc.add_paragraph()
        r = p.add_run(linea[:150])
        r.font.name = "Consolas"
        r.font.size = Pt(tam)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        p.paragraph_format.space_after = Pt(0)
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), "F2F2F2")
        p._p.get_or_add_pPr().append(sh)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------ encabezado y pie
hr = sec.header.paragraphs[0].add_run(
    "Universidad Peruana Unión · EP Ingeniería de Sistemas · DAM 2026-2 · Actividad Autónoma N.º 08")
hr.font.size = Pt(8)
hr.font.color.rgb = AZUL
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Julio Fabián Rodríguez Bazán · PharmaMobil · 01/10/2026")
fr.font.size = Pt(8)

# ------------------------------------------------------------ portada
parrafo()
parrafo()
parrafo("UNIVERSIDAD PERUANA UNIÓN", True, 22, True, AZUL)
parrafo("Facultad de Ingeniería y Arquitectura", False, 13, True)
parrafo("EP Ingeniería de Sistemas", False, 13, True, espacio=24)
parrafo("Desarrollo de Aplicaciones Móviles (DAM)", True, 14, True)
parrafo("Ciclo VI · Semestre 2026-2", False, 12, True, espacio=30)
parrafo("ACTIVIDAD AUTÓNOMA N.º 08", True, 20, True, AZUL, espacio=6)
parrafo("PRUEBAS INTEGRALES DE CRUD REST, TAXONOMÍA DE ERRORES Y PRUEBAS AUTOMATIZADAS MULTIPLATAFORMA",
        True, 14, True, AZUL, espacio=36)
ficha([
    ("Estudiante / Autor", "Julio Fabián Rodríguez Bazán"),
    ("Repositorio GitHub", "https://github.com/fabianrodriguez33/Pharma-Mobil"),
    ("Rama", "feature/ktor-client"),
    ("Fecha de entrega", "01/10/2026"),
])
salto()

# ------------------------------------------------------------ entorno
doc.add_heading("Entorno y metodología", 1)
parrafo("Todas las evidencias se obtuvieron el 01/10/2026 contra el backend PharmaSoft real (Spring Boot, Oracle Free "
        "en Docker, puerto 8080) y la aplicación PharmaMobil instalada en el emulador Pixel 8 (Android, 1080×2400), "
        "que accede al backend mediante http://10.0.2.2:8080/. Las respuestas de la matriz del Producto 1 provienen "
        "de peticiones curl reales (s08/curl_crud.log); los escenarios del Producto 2 se ejecutaron desde la "
        "aplicación y se contrastaron con el logcat de Ktor.")
parrafo("Observaciones respecto a la guía, verificadas en el servidor:", True)
for t in (
    "El DELETE de PharmaSoft es una baja lógica (estado = false). La segunda eliminación del mismo id no devuelve 404 "
    "sino 409 Conflict con el mensaje «El producto … ya se encuentra inactivo»; la app lo traduce a ErrorApi.Conflicto.",
    "El precio 0 o negativo lo detecta antes la validación local del formulario. Para provocar el 400 del servidor "
    "sobre precioError se envió 0.001, que supera la validación local pero el servidor rechaza («El precio debe ser "
    "mayor que cero»). El precio 0 enviado directamente por curl también devuelve 400.",
    "La app lista con tamanio=20 (valor por defecto de ProductoApi); la matriz usa tamanio=10 como pide la guía.",
    "El escenario 7 se reprodujo reduciendo temporalmente requestTimeoutMillis a 1 ms; el valor original (15 000 ms) "
    "fue restaurado.",
    "El escenario 3 se reprodujo borrando físicamente en la base de pruebas un producto de prueba creado por la app, "
    "de modo que la lista de la app quedó desactualizada y el PUT devolvió 404.",
):
    p = doc.add_paragraph(t, style="List Bullet")
    for r in p.runs:
        r.font.size = Pt(9.5)

# ------------------------------------------------------------ producto 1
doc.add_heading("Producto 1: Matriz de operaciones CRUD REST (PharmaSoft)", 1)
parrafo("Recurso /api/v1/productos. El código recibido proviene de s08/curl_crud.log.")
tabla(
    ["#", "Método y ruta", "Parámetros / Payload", "DTOs", "HTTP esperado", "HTTP recibido", "Resultado observado"],
    [
        ["1", "GET /api/v1/productos?pagina=0&tamanio=10", "pagina=0, tamanio=10",
         "PaginaResponseDto<ProductoResponseDto>", "200 OK", "200 OK",
         "Página con contenido, totalElementos y totalPaginas"],
        ["2", "GET /api/v1/productos/{id}", "id = 21", "ProductoResponseDto", "200 OK", "200 OK",
         "Producto 21 «Loratadina 10mg»"],
        ["3", "POST /api/v1/productos",
         '{"nombre":"Loratadina 10mg","precio":8.9,"stock":30,"estado":true,"categoriaId":1}',
         "ProductoRequestDto → ProductoResponseDto", "201 Created", "201 Created", "Producto creado con id 21"],
        ["4", "PUT /api/v1/productos/{id}",
         '{"nombre":"Loratadina 10mg Forte","precio":9.9,"stock":35,...}',
         "ProductoRequestDto → ProductoResponseDto", "200 OK", "200 OK",
         "Nombre, precio y stock actualizados; fechaModificacion asignada"],
        ["5", "DELETE /api/v1/productos/{id}", "id = 21", "Sin cuerpo", "204 No Content", "204 No Content",
         "Respuesta vacía; ProductoApi.eliminar no llama a body(), no se deserializa JSON"],
    ],
    anchos=[0.6, 3.0, 3.2, 2.7, 1.7, 1.7, 3.7], tam=8)
parrafo("Extracto del registro de peticiones (curl_crud.log):", True)
codigo("""\
### POST /api/v1/productos           -> HTTP 201
{"id":21,"nombre":"Loratadina 10mg","precio":8.9,"stock":30,"estado":true,"categoriaId":1,...}
### PUT /api/v1/productos/21         -> HTTP 200
{"id":21,"nombre":"Loratadina 10mg Forte","precio":9.9,"stock":35,...,"fechaModificacion":"2026-10-01T21:00:36.85"}
### DELETE /api/v1/productos/21      -> HTTP 204   (cuerpo vacío)""")
parrafo("Evidencias en la aplicación (emulador Pixel 8):", True)
imagen("01_listado_paginado.png", "Figura 1. GET paginado: inventario cargado desde PharmaSoft (Fase.ConProductos).")
imagen("04_crear_exitoso.png", "Figura 2. POST 201: «Amoxicilina 500mg» registrado y lista recargada.")
imagen("05_editar_formulario.png", "Figura 3. Formulario precargado para editar el producto (PUT).")
imagen("06_editar_exitoso.png", "Figura 4. PUT 200: stock actualizado de 20 a 35 unidades.")
imagen("07_eliminar_204.png", "Figura 5. DELETE 204: el producto desaparece del listado.")

# ------------------------------------------------------------ producto 2
doc.add_heading("Producto 2: Bitácora de 8 escenarios de fallo controlados", 1)
parrafo("Resumen de la taxonomía ErrorApi aplicada en cada escenario:")
tabla(
    ["Esc.", "Disparador", "HTTP / excepción", "ErrorApi", "Mensaje en la UI"],
    [
        ["1", "POST con nombre «Ab»", "400", "Validacion",
         "El nombre debe tener entre 3 y 150 caracteres (bajo nombreError)"],
        ["2", "POST con precio 0.001 (0 por curl)", "400", "Validacion",
         "El precio debe ser mayor que cero (bajo precioError)"],
        ["3", "PUT sobre id inexistente", "404", "NoEncontrado", "El producto ya no existe"],
        ["4", "DELETE repetido sobre el mismo id", "204 y luego 409", "Conflicto",
         "El producto … ya se encuentra inactivo"],
        ["5", "POST con nombre duplicado", "409", "Conflicto",
         "Ya existe un producto con el nombre Paracetamol 500mg"],
        ["6", "Modo avión", "ConnectException", "SinConexion", "Sin conexión con el servidor"],
        ["7", "requestTimeoutMillis = 1", "HttpRequestTimeoutException", "TiempoAgotado",
         "El servidor tardó demasiado en responder"],
        ["8", "Cancelar el Scope en curso", "CancellationException", "(ninguno)",
         "Sin mensaje; la corrutina se cancela limpiamente"],
    ],
    anchos=[1.0, 4.2, 3.4, 2.6, 5.4], tam=8.5)

esc = [
    ("Escenario 1: Validación del servidor (nombre corto)", "01/10/2026 · 21:16", [
        ("Pasos", "Inventario → Nombre «Ab», Precio 5, Stock 1 → Registrar."),
        ("Código HTTP", "400 Bad Request"),
        ("ErrorResponseDto", 'validationErrors = {"nombre":"El nombre debe tener entre 3 y 150 caracteres"}'),
        ("ErrorApi producido", "ErrorApi.Validacion(porCampo = {nombre: …})"),
        ("Mensaje en la UI", "«El nombre debe tener entre 3 y 150 caracteres» bajo el campo Nombre (nombreError)."),
        ("Conclusión técnica", "El 400 se mapea por campo; la Fase no cambia a Error y Operacion vuelve a Inactiva."),
    ], "02_s1_nombre_corto.png", "Figura 6. Escenario 1: error 400 bajo el campo Nombre."),
    ("Escenario 2: Precio inválido", "01/10/2026 · 21:17", [
        ("Pasos", "Nombre «Prueba Precio», Precio 0.001, Stock 1 → Registrar. (El valor 0 lo frena la validación "
                  "local; por curl también devuelve 400.)"),
        ("Código HTTP", "400 Bad Request"),
        ("ErrorResponseDto", 'validationErrors = {"precio":"El precio debe ser mayor que cero"}'),
        ("ErrorApi producido", "ErrorApi.Validacion(porCampo = {precio: …})"),
        ("Mensaje en la UI", "«El precio debe ser mayor que cero» bajo el campo Precio (precioError)."),
        ("Conclusión técnica", "El servidor es la fuente de verdad: lo que la validación local deja pasar, el 400 "
                               "lo corrige por campo."),
    ], "03_s2_precio_invalido.png", "Figura 7. Escenario 2: error 400 bajo el campo Precio."),
    ("Escenario 3: Recurso inexistente", "01/10/2026 · 21:22", [
        ("Pasos", "Crear «Temporal S3», borrarlo físicamente en la base de pruebas, editarlo en la app (lista "
                  "desactualizada) y pulsar Guardar cambios."),
        ("Código HTTP", "404 Not Found (PUT /api/v1/productos/25)"),
        ("ErrorResponseDto", 'message = "Producto no encontrado con id: …" (mismo formato que para el id 999999)'),
        ("ErrorApi producido", "ErrorApi.NoEncontrado"),
        ("Mensaje en la UI", "«El producto ya no existe»"),
        ("Conclusión técnica", "El 404 llega como Operacion.Fallida; el formulario y la lista se conservan."),
    ], "11_s3_no_encontrado_404.png", "Figura 8. Escenario 3: 404 traducido a «El producto ya no existe»."),
    ("Escenario 4: Doble eliminación", "01/10/2026 · 21:21", [
        ("Pasos", "Crear «Temporal S4»; DELETE por curl (1.ª llamada); pulsar Eliminar en la app sobre el mismo id "
                  "(2.ª llamada). Registro en s08/doble_delete.log."),
        ("Código HTTP", "1.ª: 204 No Content · 2.ª: 409 Conflict (la guía esperaba 404; el backend hace baja lógica)"),
        ("ErrorResponseDto", 'message = "El producto Temporal S4 ya se encuentra inactivo"'),
        ("ErrorApi producido", "ErrorApi.Conflicto(mensaje)"),
        ("Mensaje en la UI", "«El producto Temporal S4 ya se encuentra inactivo»"),
        ("Conclusión técnica", "Con baja lógica el segundo DELETE es una violación de regla de negocio (409), no un "
                               "recurso inexistente."),
    ], "10_s4_doble_eliminacion.png", "Figura 9. Escenario 4: segunda eliminación rechazada con 409."),
    ("Escenario 5: Regla de negocio / conflicto", "01/10/2026 · 21:20", [
        ("Pasos", "Registrar un producto con nombre duplicado «Paracetamol 500mg»."),
        ("Código HTTP", "409 Conflict"),
        ("ErrorResponseDto", 'message = "Ya existe un producto con el nombre Paracetamol 500mg"'),
        ("ErrorApi producido", "ErrorApi.Conflicto(mensaje)"),
        ("Mensaje en la UI", "«Ya existe un producto con el nombre Paracetamol 500mg»"),
        ("Conclusión técnica", "El mensaje del servidor se muestra tal cual; no se pierde información de negocio."),
    ], "08_s5_conflicto_409.png", "Figura 10. Escenario 5: conflicto 409 con el mensaje emergente del servidor."),
    ("Escenario 6: Servidor caído / sin conexión", "01/10/2026 · 21:22", [
        ("Pasos", "Activar Modo Avión en el emulador y pulsar Eliminar sobre un producto."),
        ("Excepción", "java.net.ConnectException: Failed to connect to /10.0.2.2:8080 (logcat_ktor_crud.log)"),
        ("ErrorApi producido", "ErrorApi.SinConexion"),
        ("Mensaje en la UI", "«Sin conexión con el servidor» como Operacion.Fallida. La carga inicial muestra la "
                             "pantalla de error con ícono de nube tachada y botón Reintentar (ver Figura 12)."),
        ("Conclusión técnica", "La excepción de red se traduce sin exponer tipos de Ktor a la presentación."),
    ], "12_s6_modo_avion.png", "Figura 11. Escenario 6: Modo Avión activo (ícono en la barra de estado)."),
    ("Escenario 7: Tiempo de espera agotado", "01/10/2026 · 21:23", [
        ("Pasos", "requestTimeoutMillis = 1 en HttpClientFactory, reinstalar y abrir Inventario. Valor original "
                  "restaurado después."),
        ("Excepción", "HttpRequestTimeoutException: Request timeout has expired (request_timeout=1 ms), "
                      "ver logcat_timeout.log"),
        ("ErrorApi producido", "ErrorApi.TiempoAgotado"),
        ("Mensaje en la UI", "«El servidor tardó demasiado en responder» con botón Reintentar."),
        ("Conclusión técnica", "El timeout se clasifica antes del caso genérico y no se confunde con una "
                               "cancelación de corrutina."),
    ], "13_s7_timeout.png", "Figura 12. Escenario 7: pantalla de error con ícono de red y botón Reintentar."),
    ("Escenario 8: Cancelación de corrutina", "01/10/2026 · prueba automatizada", [
        ("Pasos", "Pruebas ErrorApiTest en commonTest: toErrorApi() con CancellationException y cancelación del Job "
                  "mientras traducirErrores {} espera la respuesta."),
        ("Excepción", "CancellationException"),
        ("ErrorApi producido", "Ninguno: la excepción se relanza antes del catch genérico."),
        ("Mensaje en la UI", "Ninguno."),
        ("Conclusión técnica", "Verificado por pruebas (4/4 en ErrorApiTest): la corrutina cancelada termina "
                               "limpiamente, no se convierte en ErrorApi.Desconocido y no deja la UI en error."),
    ], None, None),
]
for titulo, fecha, pares, img, leyenda in esc:
    doc.add_heading(titulo, 2)
    ficha([("Fecha/Hora", fecha), ("Plataforma", "Android (emulador Pixel 8) · backend PharmaSoft :8080")] + pares)
    if img:
        imagen(img, leyenda)

# ------------------------------------------------------------ producto 3
doc.add_heading("Producto 3: Pruebas automatizadas multiplataforma (commonTest)", 1)
parrafo("Comando ejecutado: ./gradlew testAndroidHostTest assembleDebug → BUILD SUCCESSFUL (s08/gradle_build.log).",
        True)
res = []
tot = fall = 0
for ruta in sorted(glob.glob(os.path.join(AQUI, "..", "..", "shared", "build", "test-results", "*", "*.xml"))):
    s = open(ruta, encoding="utf-8").read()
    m = re.search(r'<testsuite name="([^"]+)" tests="(\d+)" skipped="\d+" failures="(\d+)" errors="(\d+)"', s)
    tot += int(m[2])
    fall += int(m[3]) + int(m[4])
    res.append([m[1].replace("pe.edu.upeu.pharmamobil.", ""), m[2], str(int(m[3]) + int(m[4]))])
tabla(["Requerimiento", "Prueba", "Qué verifica"], [
    ["Prueba 1: carga exitosa", "ProductoViewModelTest.laCargaPasaDeCargandoAConProductos",
     "Con la lista retenida la fase es Fase.Cargando; al liberarla pasa a Fase.ConProductos con la lista esperada."],
    ["Prueba 2: listado vacío", "ProductoViewModelTest.arrancaEnSinProductosCuandoElInventarioEstaVacio",
     "Con repositorio vacío la fase termina en Fase.SinProductos."],
    ["Prueba 3: error de validación",
     "losErroresDeValidacionCaenEnElFormularioNoEnLaFase y elErrorDeValidacionDelServidorCaeBajoCadaCampo",
     "Los mensajes quedan en nombreError, precioError o stockError sin cambiar la fase a Fase.Error."],
    ["Prueba 4: eliminación", "ProductoViewModelTest.eliminarPasaPorEnCursoEliminarYVuelveAInactiva",
     "Durante eliminar() el estado es Operacion.EnCurso(Tipo.Eliminar); al terminar vuelve a Operacion.Inactiva."],
    ["Escenario 8", "ErrorApiTest (4 pruebas)",
     "CancellationException se relanza; cancelar el Job no produce ErrorApi; conexión fallida → SinConexion; "
     "ConnectTimeoutException → TiempoAgotado."],
], anchos=[3.4, 5.6, 7.6], tam=8.5)
parrafo("Resultados por clase de prueba (XML de Gradle):", True)
tabla(["Clase", "Pruebas", "Fallos"], res + [["TOTAL", str(tot), str(fall)]], anchos=[11, 2.5, 2.5], tam=8.5)
parrafo("Salida de la ejecución:", True)
log = open(os.path.join(EV, "gradle_build.log"), encoding="utf-8", errors="ignore").read().splitlines()
codigo("\n".join([l for l in log if re.search(r"testAndroidHostTest|assembleDebug|BUILD|actionable", l)][-8:]))

# ------------------------------------------------------------ producto 4
doc.add_heading("Producto 4: Evidencias de ejecución y registro de Ktor", 1)
parrafo("Registro del cliente Ktor (logcat filtrado por «HttpClient»; s08/logcat_ktor_crud.log):", True)
kt = open(os.path.join(EV, "logcat_ktor_crud.log"), encoding="utf-8", errors="ignore").read().splitlines()
kt = [re.sub(r"^\S+ \S+\s+\d+\s+\d+ I System.out: ", "", l) for l in kt if "RESPONSE" in l or "REQUEST" in l]
codigo("\n".join(kt[:40]))
parrafo("Registro de la excepción de timeout y de conexión:", True)
codigo(re.sub(r"\S+ \S+\s+\d+\s+\d+ I System.out: ", "",
              open(os.path.join(EV, "logcat_timeout.log"), encoding="utf-8", errors="ignore").read()))
parrafo("README.md actualizado con las secciones «CRUD REST PharmaSoft» y «Manejo de Errores».")

salida = os.path.join(AQUI, "S08_ActividadAutonoma_RodriguezBazan.docx")
doc.save(salida)
print(salida, os.path.getsize(salida))
