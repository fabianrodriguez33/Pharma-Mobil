"""Genera docs/entrega/S11_PersistenciaOffline_RodriguezBazan.docx con python-docx."""
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, "..", "entrega", "S11_PersistenciaOffline_RodriguezBazan.docx")
NAVY = RGBColor(0x1F, 0x38, 0x64)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for nombre, size in (("Heading 1", 17), ("Heading 2", 13.5)):
    h = doc.styles[nombre]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = NAVY


def p(texto, negrita=False, cursiva=False):
    par = doc.add_paragraph()
    r = par.add_run(texto)
    r.bold, r.italic = negrita, cursiva
    return par


def tabla(filas, anchos):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Light Grid Accent 1"
    for i, fila in enumerate(filas):
        for j, celda in enumerate(fila):
            c = t.cell(i, j)
            c.width = Cm(anchos[j])
            c.text = celda
            for run in c.paragraphs[0].runs:
                run.font.size = Pt(9.5)
                run.bold = i == 0
    doc.add_paragraph()


def foto(archivo, pie, ancho=6.0):
    doc.add_picture(os.path.join(AQUI, archivo), width=Cm(ancho))
    doc.paragraphs[-1].alignment = 1
    par = p(pie, cursiva=True)
    par.alignment = 1


t = doc.add_heading("Guía Práctica N.º 11 — Persistencia local offline-first", 0)
p("Estudiante: Julio Fabián Rodríguez Bazán\nAplicación: PharmaMobil (KMP) · Backend: PharmaSoft\n"
  "Rama: feature/offline-first-rodriguez · Fecha de pruebas: 21/10/2026")

doc.add_heading("1. Resumen", 1)
p("PharmaMobil pasó a ser offline-first: una base SQLite gestionada con SQLDelight es la única fuente de verdad "
  "de las pantallas de productos. Los repositorios exponen un Flow que las pantallas observan, y la red queda como "
  "un proceso de sincronización en dos sentidos: primero se suben los productos con pendiente = 1 (POST) y después "
  "se descarga el catálogo y se reemplazan las filas sincronizadas dentro de una transacción.")

doc.add_heading("2. Implementación", 1)
tabla([
    ["Pieza", "Archivo", "Función"],
    ["Esquema", "commonMain/sqldelight/.../db/Producto.sq", "Tabla producto con columna pendiente y consultas"],
    ["Drivers", "CreadorDeDriver (+Android, +Ios)", "AndroidSqliteDriver, NativeSqliteDriver; JdbcSqliteDriver en pruebas"],
    ["Capa local", "data/local/ProductoLocal.kt", "observar(): Flow, guardarPendiente (id negativo), reemplazarSincronizados (transacción)"],
    ["Repositorio", "data/repository/ProductoRepositorioOffline.kt", "registrar() local + intento de sync; sincronizar(): sube pendientes y luego descarga paginado"],
    ["Contrato", "domain/repository/ProductoRepository.kt", "observar(): Flow y sincronizar()"],
    ["ViewModels", "ProductosViewModel, ProductoViewModel", "Se suscriben al Flow en init; sin recargas manuales"],
    ["DI", "AppModule, NetworkModule, PlatformModule", "PharmaDb, ProductoLocal y driver por plataforma"],
], [2.6, 6.8, 7.2])

doc.add_heading("3. Pruebas automatizadas (commonTest)", 1)
p("Comando: ./gradlew testAndroidHostTest assembleDebug → BUILD SUCCESSFUL. Suite completa del proyecto: 70 pruebas, 0 fallos.")
tabla([
    ["Prueba (ProductoRepositorioOfflineTest)", "Verifica", "Resultado"],
    ["loRegistradoSinRedQuedaConIdNegativoYPendiente", "Sin red: id < 0 y pendiente = 1; observar() lo emite", "OK"],
    ["sincronizarSubePendientesYReemplazaLaListaLocal", "Orden POST → GET; 0 pendientes; lista = servidor", "OK"],
    ["conElServidorFueraDeLineaObservarEmiteLosDatosGuardados", "Con 503, observar() entrega la base local sin crash", "OK"],
    ["elViewModelMuestraLoRegistradoSinRecargarAMano", "El ViewModel refleja el alta vía Flow", "OK"],
], [7.4, 7.0, 2.2])

doc.add_heading("4. Prueba en emulador Android (Pixel 8, API 37.1)", 1)
p("Ciclo completo contra PharmaSoft real (Oracle, Spring Boot en :8080), usuario admin del seed del proyecto.")
foto("01_catalogo_con_servidor.png", "Figura 1. Catálogo con el servidor encendido (3 productos).")
foto("02_catalogo_sin_red.png", "Figura 2. Backend detenido y modo avión: la app abre y lista desde pharma.db, con el aviso «Sin conexión».")
foto("03_registro_sin_red.png", "Figura 3. Producto «Ibuprofeno_Offline» registrado sin red: aparece de inmediato en el inventario (4 productos).")
foto("04_tras_reconectar.png", "Figura 4. Tras reconectar y reabrir el catálogo: el aviso desapareció y el producto sigue listado.")
p("Estado en PharmaSoft tras la sincronización (GET /api/v1/productos, productos activos):", negrita=True)
with open(os.path.join(AQUI, "05_servidor_tras_sync.txt"), encoding="utf-8") as f:
    par = doc.add_paragraph()
    r = par.add_run("id  nombre  precio  stock  estado\n" + f.read())
    r.font.name, r.font.size = "Consolas", Pt(9.5)
p("El producto creado sin red llegó al servidor con id real 47 (el id local negativo se descartó), sin duplicados.")

doc.add_heading("5. Errores hallados durante la prueba real y corregidos", 1)
tabla([
    ["Hallazgo", "Causa", "Corrección"],
    ["sincronizar() fallaba con «El tamaño de página debe estar entre 1 y 100»", "Se pedía una página de 1000; el backend limita a 100", "Descarga paginada de 100 en 100 hasta la última página"],
    ["La app se cerraba (NullPointerException) al abrir el Catálogo", "productos se declaraba después del init del ViewModel y estaba en null cuando sincronizar() lo usaba", "Declarar el estado antes del init"],
], [5.6, 5.8, 5.2])

doc.add_heading("6. Lista de cotejo", 1)
tabla([
    ["N.º", "Criterio", "Cumple", "Evidencia"],
    ["1", "El esquema SQLDelight compila y genera PharmaDb", "Sí", "Build exitoso"],
    ["2", "Driver en Android, iOS y pruebas", "Sí (iOS sin ejecutar)", "CreadorDeDriver*, JdbcSqliteDriver"],
    ["3", "La capa local no importa Ktor ni DTO", "Sí", "ProductoLocal.kt"],
    ["4", "Repositorio expone observar(): Flow y sincronizar()", "Sí", "ProductoRepository"],
    ["5", "El ViewModel se suscribe al Flow", "Sí", "Ambos ViewModels; test 4"],
    ["6", "Lista productos con backend apagado", "Sí", "Figura 2"],
    ["7", "Creado sin red queda pendiente", "Sí", "Test 1; Figura 3"],
    ["8", "sincronizar() sube antes de bajar", "Sí", "Test 2; código"],
    ["9", "Aparece en PharmaSoft tras reconectar", "Sí", "Figura 4; id 47 en el servidor"],
    ["10", "Al menos 3 pruebas en commonTest, todas pasan", "Sí", "4 pruebas, 70 en total"],
], [1.0, 7.0, 3.6, 5.0])

doc.add_heading("7. Limitaciones", 1)
p("• iOS: el entorno de desarrollo es Windows; el código de iOS (NativeSqliteDriver y registro en Koin) está implementado "
  "pero no se compiló ni ejecutó, por lo que no hay captura en iOS.\n"
  "• La imagen de un producto creado sin red no puede subirse hasta que el producto se sincronice (la foto se sube con el id real).")

doc.save(SALIDA)
print(SALIDA)
