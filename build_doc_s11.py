"""Genera docs/entrega/S11_ActividadAutonoma_RodriguezBazan.docx con python-docx."""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.abspath(__file__))
EV = os.path.join(RAIZ, "docs", "actividad11")
SALIDA = os.path.join(RAIZ, "docs", "entrega", "S11_ActividadAutonoma_RodriguezBazan.docx")
NAVY = RGBColor(0x1F, 0x38, 0x64)
CELESTE = "9DC3E6"
ALT = "EAF3FB"
HEAD = "1F3864"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for nombre, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
    h = doc.styles[nombre]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def p(texto="", negrita=False, cursiva=False, size=None, align=None, after=6, color=None):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(after)
    if align is not None:
        par.alignment = align
    r = par.add_run(texto)
    r.bold, r.italic = negrita, cursiva
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return par


def vineta(texto):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_after = Pt(2)
    par.add_run(texto)


def mono(texto, size=8):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    r = par.add_run(texto)
    r.font.name, r.font.size = "Consolas", Pt(size)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def sombrear(celda, color):
    pr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    pr.append(sh)


def bordes(tabla, color=CELESTE):
    tbl = tabla._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    pr.append(b)


def tabla(filas, anchos, size=9.5):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes(t)
    for i, fila in enumerate(filas):
        for j, texto in enumerate(fila):
            c = t.cell(i, j)
            c.width = Cm(anchos[j])
            c.text = ""
            run = c.paragraphs[0].add_run(texto)
            run.font.size = Pt(size)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                sombrear(c, HEAD)
            elif i % 2 == 0:
                sombrear(c, ALT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def fotos(items, ancho=5.2):
    """items: lista de (archivo, pie); se colocan en una fila de tabla sin bordes."""
    t = doc.add_table(rows=2, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (archivo, pie) in enumerate(items):
        c = t.cell(0, j)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].add_run().add_picture(os.path.join(EV, archivo), width=Cm(ancho))
        d = t.cell(1, j)
        d.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = d.paragraphs[0].add_run(pie)
        r.italic, r.font.size = True, Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def texto_de(archivo):
    with open(os.path.join(EV, archivo), encoding="utf-8") as f:
        return f.read().rstrip("\n")


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ───────────────────────── PORTADA ─────────────────────────
for _ in range(2):
    p()
p("UNIVERSIDAD PERUANA UNIÓN", negrita=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=4)
p("Facultad de Ingeniería y Arquitectura · EP Ingeniería de Sistemas", size=13,
  align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=24)
p("Asignatura: Desarrollo de Aplicaciones Móviles (DAM) · Ciclo VI / Semestre 2026-2", size=12,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
p("ACTIVIDAD AUTÓNOMA N.º 11", negrita=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=8)
p("PRUEBAS UNITARIAS EN COMMONTEST Y EVIDENCIA DE ESCENARIOS OFFLINE-FIRST EN PHARMAMOBIL",
  negrita=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=50)
tabla([
    ["Estudiante", "Julio Fabián Rodríguez Bazán"],
    ["Repositorio GitHub", "https://github.com/fabianrodriguez33/Pharma-Mobil"],
    ["Rama", "feature/offline-first-rodriguez"],
    ["Fecha de entrega", "26/10/2026"],
], [4.5, 11.5], size=11)
salto()

# ───────────────────────── 1. RESUMEN ─────────────────────────
doc.add_heading("1. Resumen ejecutivo", 1)
p("Esta actividad cierra la Unidad 2: valida con pruebas automatizadas y con escenarios reales en emulador que "
  "PharmaMobil es offline-first, agrega un indicador de la última sincronización y prepara la sustentación del "
  "Producto U2. Todo se probó contra PharmaSoft real (Spring Boot + Oracle) y un emulador Pixel 8 (Android API 37.1).")
tabla([
    ["Producto", "Entregable", "Estado"],
    ["1 · Pruebas unitarias (7 pts)", "6 pruebas del enunciado + 2 de apoyo en ProductoRepositorioOfflineTest; 80 pruebas en el proyecto, 0 fallos", "Completo"],
    ["2 · Bitácora offline (5 pts)", "5 escenarios ejecutados en el emulador con capturas, pasos y resultado", "Completo"],
    ["3 · Indicador de sincronización (4 pts)", "Marca de tiempo persistida + texto en el encabezado + botón «Sincronizar»", "Completo"],
    ["4 · Guion de sustentación (4 pts)", "Guion de 5 criterios, 8 minutos en total", "Completo"],
], [4.6, 9.2, 2.2])
p("Hallazgos técnicos durante la ejecución (corregidos antes de la entrega):", negrita=True)
vineta("Crash al abrir el catálogo: el estado del ViewModel se declaraba después del bloque init (NullPointerException).")
vineta("El backend limita la página a 100 productos; la descarga ahora recorre todas las páginas.")
vineta("La marca de sincronización exigió una migración SQLDelight (1.sqm, esquema v2); se verificó en el emulador "
       "que una base v1 ya instalada migra sin perder datos.")

# ───────────────────────── 2. PRODUCTO 1 ─────────────────────────
doc.add_heading("2. Producto 1 — Suite de pruebas unitarias en commonTest", 1)
p("Archivo: shared/src/commonTest/kotlin/pe/edu/upeu/pharmamobil/ProductoRepositorioOfflineTest.kt. "
  "Cada prueba crea su propia base con JdbcSqliteDriver(IN_MEMORY) y simula el servidor con MockEngine de Ktor, "
  "por lo que no comparten estado.")
tabla([
    ["N.º", "Prueba", "Qué verifica", "Resultado"],
    ["1", "prueba1_sincronizarConRedGuardaLosRemotos…", "Con 200 OK la base guarda los productos remotos y observar() emite la lista completa", "Pasa"],
    ["2", "prueba2_registrarSinRedGuardaIdNegativo…", "Con 503 el producto se guarda con id < 0 y pendiente = 1; el Flow lo emite", "Pasa"],
    ["3", "prueba3_alReconectarSubeElPendiente…", "POST antes del GET; la fila queda con pendiente = 0 y el id definitivo del servidor", "Pasa"],
    ["4", "prueba4_laBajadaNoBorraLoPendiente", "reemplazarSincronizados() borra solo pendiente = 0 y conserva lo creado sin red", "Pasa"],
    ["5", "prueba5_idsNegativosYPositivosConviven", "Ids -1, -2 conviven con 1, 2 del servidor sin sobrescribirse", "Pasa"],
    ["6", "prueba6_sincronizarDosVecesNoDuplica…", "Dos sincronizaciones seguidas no duplican filas", "Pasa"],
    ["+", "conElServidorFueraDeLinea…", "Con el servidor caído, observar() emite lo guardado sin crash", "Pasa"],
    ["+", "laSincronizacionExitosaGuardaSuMarca…", "La marca de tiempo solo se guarda si la sincronización tuvo éxito", "Pasa"],
], [1.0, 5.6, 7.4, 2.0], size=9)
p("Pruebas complementarias: TextoSincronizacionTest (5 pruebas del texto del indicador) y "
  "ProductoViewModelOfflineTest (el ViewModel refleja un alta sin recargar a mano).")
doc.add_heading("Resultado por clase (extraído de los reportes JUnit)", 3)
mono(texto_de("resultados_tests.txt"), size=7.5)
doc.add_paragraph()
doc.add_heading("Salida completa de ./gradlew testAndroidHostTest", 3)
p("Comando ejecutado en D:\\DAM\\pharmaMobil-master con --rerun-tasks (BUILD SUCCESSFUL). El registro HTTP que se ve "
  "corresponde a las peticiones simuladas por MockEngine.", cursiva=True)
for linea in texto_de("gradle_test_output.txt").splitlines():
    mono(linea[:150], size=6.5)
salto()

# ───────────────────────── 3. PRODUCTO 2 ─────────────────────────
doc.add_heading("3. Producto 2 — Bitácora de 5 escenarios offline (emulador Pixel 8)", 1)
p("Entorno: emulador Pixel 8 (Android API 37.1), PharmaSoft en localhost:8080 (Oracle), usuario admin del seed del "
  "proyecto. Modo Avión activado con «adb shell cmd connectivity airplane-mode». La base local se inspeccionó "
  "extrayendo pharma.db con «adb exec-out run-as … cat databases/pharma.db».")

doc.add_heading("Escenario 1 — Apertura sin red con datos guardados", 2)
p("Pasos:", negrita=True)
vineta("Con datos ya sincronizados, activar Modo Avión y cerrar la app por completo (force-stop).")
vineta("Abrir PharmaMobil y entrar a Catálogo.")
p("Resultado observado: el catálogo se lista desde SQLite (3 productos) y aparece el aviso «Sin conexión: mostrando "
  "datos guardados» junto con la antigüedad de los datos. No hay pantalla en blanco ni error.", negrita=False)
fotos([("s1_apertura_sin_red.png", "Figura 1. Catálogo en Modo Avión con aviso y «Actualizado hace 12 minutos».")])

doc.add_heading("Escenario 2 — Registro en Modo Avión", 2)
p("Pasos:", negrita=True)
vineta("Con Modo Avión activo, ir a Inventario y registrar «Amoxicilina_Offline», S/ 6.80, 30 unidades.")
p("Resultado observado: aparece de inmediato en la lista (4 productos) sin recargar. En SQLite quedó con id negativo "
  "y pendiente = 1:")
mono(texto_de("s2_base_local.txt"), size=9)
doc.add_paragraph()
fotos([("s2_registro_modo_avion.png", "Figura 2. Producto registrado sin red y visible al instante.")])

doc.add_heading("Escenario 3 — Reconexión y sincronización exitosa", 2)
p("Pasos:", negrita=True)
vineta("Antes de reconectar se editó el pendiente a S/ 7.20 (ver Escenario 5).")
vineta("Desactivar Modo Avión y abrir Catálogo: la app sincroniza al entrar; se pulsó además «Sincronizar».")
vineta("Verificar en SQLite y en Swagger UI (GET /api/v1/productos).")
p("Resultado observado: el POST se envió, la fila -1 pasó a id 48 con pendiente = 0 y conservó el precio editado "
  "offline (7.20). Swagger UI muestra el producto con id 48 y estado true.")
mono(texto_de("s3_base_local_tras_sync.txt"), size=9)
doc.add_paragraph()
fotos([
    ("s3b_tras_sincronizar.png", "Figura 3. Tras reconectar: «Actualizado hace un momento»."),
    ("s3c_swagger.jpg", "Figura 4. Swagger UI: GET /api/v1/productos (200) con id 48."),
], ancho=7.0)

doc.add_heading("Escenario 4 — Sincronización con token vencido", 2)
p("Para provocar un vencimiento real se reinició PharmaSoft con app.jwt.expiracion-minutos = 1 (el login devolvió "
  "expiraEnSegundos = 60).", cursiva=True)
p("Pasos:", negrita=True)
vineta("Iniciar sesión (token de 60 s), activar Modo Avión y registrar «Token_Vencido» (queda pendiente).")
vineta("Esperar más de 60 s para que el JWT venza; desactivar Modo Avión y abrir Catálogo (sincroniza).")
vineta("Observar el 401, comprobar la base y volver a iniciar sesión.")
p("Resultado observado: el servidor respondió 401 (registrado en Logcat), la app cerró la sesión y mostró «Tu sesión "
  "expiró. Inicia sesión nuevamente». El trabajo pendiente NO se perdió; tras reingresar se envió y la fila -1 pasó "
  "a id 49 con pendiente = 0.")
tabla([
    ["Momento", "Contenido de la tabla producto (id | nombre | pendiente)"],
    ["Antes del 401", "-1 | Token_Vencido | 1 (pendiente, creado en Modo Avión)"],
    ["Tras el 401 (sesión cerrada)", "-1 | Token_Vencido | 1 → intacto"],
    ["Tras reingresar y sincronizar", "49 | Token_Vencido | 0 → enviado"],
], [5.0, 11.0])
fotos([
    ("s4a_sesion_vencida.png", "Figura 5. Aviso de sesión vencida tras el 401."),
    ("s4b_tras_relogin.png", "Figura 6. Tras reingresar: pendiente enviado y catálogo actualizado."),
])
p("Diseño: el cierre de sesión solo borra el token y el perfil (AlmacenSeguro); la base pharma.db no se toca, por eso "
  "lo pendiente sobrevive. Además, si el 401 llegara a la pantalla de catálogo el aviso diría «Sesión vencida… tus "
  "cambios pendientes siguen guardados».")

doc.add_heading("Escenario 5 — Edición y eliminación sin red", 2)
p("Pasos y resultados (Modo Avión activo):", negrita=True)
tabla([
    ["Acción", "Resultado observado"],
    ["Eliminar un producto ya sincronizado (Almacen 001831)", "Mensaje «Sin conexión con el servidor»; el producto permanece; nada cambia"],
    ["Editar un producto ya sincronizado (S/ 3.00 → 3.50)", "Mismo mensaje; el precio sigue en S/ 3.00"],
    ["Editar un producto pendiente (Amoxicilina_Offline 6.80 → 7.20)", "Se aplica al instante en local; al reconectar viaja en el POST (el servidor guardó 7.20)"],
], [7.0, 9.0])
fotos([
    ("s5a_eliminar_sincronizado_sin_red.png", "Figura 7. Eliminar un registro sincronizado sin red."),
    ("s5c_editar_sincronizado_sin_red.png", "Figura 8. Editar un registro sincronizado sin red."),
    ("s5b_editar_pendiente_sin_red.png", "Figura 9. Editar un pendiente: se aplica en local (S/ 7.20)."),
], ancho=4.9)
p("Argumentación técnica de la estrategia:", negrita=True)
vineta("Registros pendientes (id negativo): solo existen en el dispositivo, así que editarlos o eliminarlos es una "
       "operación local sobre una fila con pendiente = 1; el POST posterior ya lleva el estado final. No hay conflicto posible.")
vineta("Registros ya sincronizados: la app aplica write-through. La modificación o la baja se envían primero al "
       "servidor y solo si este confirma se refleja en SQLite. Sin red se rechaza con un mensaje claro y la base no cambia.")
vineta("¿Es correcto? Es seguro y consistente: no hay datos perdidos, ni estados que diverjan de PharmaSoft, y el usuario "
       "recibe retroalimentación. Pero no es offline-first completo para edición/baja: el usuario no puede trabajar esas "
       "operaciones sin red. Es una limitación consciente, no un defecto oculto.")
vineta("Mejora recomendada (fuera del alcance): una tabla outbox de operaciones pendientes (CREATE/UPDATE/DELETE) con "
       "lápidas para las bajas y resolución de conflictos por fechaModificacion/versión, ejecutada con el mismo orden "
       "«subir antes de bajar».")
vineta("Nota de verificación: la eliminación de un pendiente (borrado local de la fila con id < 0) está implementada y "
       "cubierta por el diseño del repositorio, pero no se ejercitó en el emulador.")
salto()

# ───────────────────────── 4. PRODUCTO 3 ─────────────────────────
doc.add_heading("4. Producto 3 — Indicador de última sincronización", 1)
doc.add_heading("Implementación", 2)
vineta("Persistencia: tabla sincronizacion(clave, valor) en la misma base SQLite con la marca (epoch millis, Long). "
       "Se guarda únicamente al terminar con éxito subir pendientes y descargar el catálogo. Se eligió la base y no "
       "AlmacenSeguro porque queda junto a los datos que describe, es observable con Flow y no exige cambiar la interfaz "
       "de la plataforma.")
vineta("Migración: SQLDelight 1.sqm (esquema v1 → v2). Verificado en el emulador: una base v1 ya instalada migró sin "
       "perder los productos y mostró «Nunca sincronizado».")
vineta("UI: el encabezado de ProductosScreen muestra el texto del indicador y el botón «Sincronizar» (deshabilitado "
       "mientras sincroniza). El texto se refresca cada 30 s.")
vineta("Reglas del texto: sin marca → «Nunca sincronizado»; < 1 min → «Actualizado hace un momento»; < 60 min → "
       "«Actualizado hace X minutos»; mismo día → «Actualizado hoy a las HH:mm»; otro día → «Actualizado el dd/MM a las HH:mm».")
doc.add_heading("Evidencia", 2)
p("Para los estados con antigüedad se modificó la marca en pharma.db con la app cerrada y en Modo Avión (para que la app no "
  "re-sincronizara al abrir).", cursiva=True)
fotos([
    ("e0_nunca_sincronizado.png", "Figura 10. Estado inicial: «Nunca sincronizado»."),
    ("e1_actualizado_reciente.png", "Figura 11. Datos recientes: «hace un momento»."),
    ("e2_hace_12_minutos.png", "Figura 12. «Actualizado hace 12 minutos»."),
], ancho=4.9)
fotos([
    ("e3_hoy_a_las.png", "Figura 13. «Actualizado hoy a las 17:19» (3 h)."),
    ("e4_fecha_antigua.png", "Figura 14. «Actualizado el 20/10 a las 18:20» (más de 1 día)."),
], ancho=4.9)
salto()

# ───────────────────────── 5. PRODUCTO 4 ─────────────────────────
doc.add_heading("5. Producto 4 — Guion de sustentación del Producto U2 (27 de octubre)", 1)
p("Duración total: 8 minutos (1,5 + 1,5 + 1,5 + 1,5 + 2). Preparación previa: PharmaSoft y emulador encendidos, "
  "sesión cerrada, Swagger UI abierto, y datos sincronizados al menos una vez.")


def criterio(titulo, tiempo, mostrar, decir):
    doc.add_heading(f"{titulo}  ({tiempo})", 2)
    p("Mostrar:", negrita=True, after=2)
    for m in mostrar:
        vineta(m)
    p("Decir:", negrita=True, after=2)
    for d in decir:
        vineta(d)


criterio("Criterio 1 — Consumo REST con Ktor (Sesión 7)", "1,5 min",
         ["HttpClientFactory.kt: un solo HttpClient con ContentNegotiation (JSON) y Logging con el token oculto.",
          "Motores por plataforma en PlatformModule: OkHttp (Android) y Darwin (iOS).",
          "ProductoDto/ProductoRequestDto y el log del GET /api/v1/productos con 200."],
         ["El motor se inyecta con Koin: la lógica compartida no sabe qué motor usa.",
          "Los DTO viven en data/remote y el dominio no los conoce (mapper)."])
criterio("Criterio 2 — CRUD completo (Sesión 8)", "1,5 min",
         ["Swagger UI o la app: GET lista, GET por id, POST, PUT y DELETE sobre /api/v1/productos.",
          "DELETE devuelve 204 No Content: ProductoApi.eliminar no intenta leer el cuerpo.",
          "Registrar, editar y eliminar un producto desde Inventario."],
         ["Los cinco verbos están implementados y probados.",
          "El backend hace baja lógica (estado = false); el mapper filtra los inactivos."])
criterio("Criterio 3 — Estados y errores (Sesión 8)", "1,5 min",
         ["ProductoUiState: Fase (Cargando, SinProductos, ConProductos, Error) separada de Operacion "
          "(Inactiva, EnCurso, Fallida).",
          "Enviar el formulario vacío: el 400 se traduce a errores bajo cada campo (ErrorApi.Validacion)."],
         ["Separar Fase de Operacion evita que una operación fallida borre la lista.",
          "ErrorApi es el único tipo de error que cruza a presentation; no expone Ktor."])
criterio("Criterio 4 — Autenticación JWT (Sesión 10)", "1,5 min",
         ["Login con admin; token guardado cifrado en AlmacenSeguro (AES-GCM / Keychain).",
          "Plugin Auth de Ktor agrega la cabecera Bearer; el log muestra «Authorization: ***».",
          "Roles: admin ve todo; vendedor sin formulario de productos; almacen con catálogo escritura.",
          "Un 401 devuelve al login con aviso."],
         ["El token nunca se imprime ni se guarda en claro.",
          "Un 403 mantiene la sesión; un 401 la cierra."])
criterio("Criterio 5 — Persistencia offline-first (Sesión 11)", "2 min",
         ["Modo Avión → abrir Catálogo: lista desde SQLite con «Sin conexión…» y «Actualizado hace X».",
          "Registrar un producto sin red: aparece al instante (id negativo, pendiente = 1).",
          "Reconectar y pulsar «Sincronizar»: llega a Swagger con id real.",
          "Código: ProductoRepositorioOffline.sincronizar() — subirPendientes() antes de descargar."],
         ["SQLite es la única fuente de verdad; la pantalla solo observa un Flow.",
          "El orden es sagrado: subir antes de bajar; si se invierte, la descarga borraría lo creado sin red.",
          "Los ids negativos evitan colisiones con los del servidor.",
          "Limitación conocida: editar/eliminar lo ya sincronizado requiere red (propuesta: outbox)."])
doc.add_heading("Plan de contingencia", 2)
vineta("Si el backend no responde: usar la grabación/capturas del informe y mostrar las pruebas pasando (./gradlew testAndroidHostTest).")
vineta("Si el emulador va lento: ejecutar el ciclo con las pruebas 2, 3 y 4 de ProductoRepositorioOfflineTest.")
vineta("Pregunta probable: «¿Por qué no usar AlmacenSeguro para la fecha?» → la marca vive junto a los datos que describe.")

# ───────────────────────── 6. CONCLUSIONES ─────────────────────────
doc.add_heading("6. Conclusiones y limitaciones", 1)
vineta("PharmaMobil abre y lista el catálogo sin red, registra offline, y sincroniza en el orden correcto con id definitivo.")
vineta("Las 6 pruebas del enunciado y la suite completa (80) pasan; el build termina en BUILD SUCCESSFUL.")
vineta("Limitaciones: edición/baja de registros ya sincronizados sin red; iOS no se compiló ni ejecutó (entorno Windows); "
       "la foto de un producto creado sin red solo puede subirse tras sincronizar.")

doc.save(SALIDA)
print(SALIDA)
